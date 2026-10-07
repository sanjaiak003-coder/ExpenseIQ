import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

doc = Document()

# Set page margins: Top 1", Bottom 1", Left 1.25", Right 1"
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)

    sectPr = section._sectPr
    pgBorders = parse_xml(r"""
        <w:pgBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:offsetFrom="page">
            <w:top w:val="double" w:sz="12" w:space="24" w:color="0F172A"/>
            <w:left w:val="double" w:sz="12" w:space="24" w:color="0F172A"/>
            <w:bottom w:val="double" w:sz="12" w:space="24" w:color="0F172A"/>
            <w:right w:val="double" w:sz="12" w:space="24" w:color="0F172A"/>
        </w:pgBorders>
    """)
    sectPr.append(pgBorders)

    footer = section.footer
    footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.0))
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_l = footer_table.cell(0, 0)
    p_l = cell_l.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_l = p_l.add_run("ExpenseIQ: Personal Expense and Income Tracking System")
    r_l.font.name = "Times New Roman"
    r_l.font.size = Pt(9)
    r_l.font.color.rgb = RGBColor(71, 85, 105)

    cell_r = footer_table.cell(0, 1)
    p_r = cell_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_r = p_r.add_run("Page ")
    r_r.font.name = "Times New Roman"
    r_r.font.size = Pt(9)
    r_r.font.color.rgb = RGBColor(71, 85, 105)
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    p_r._p.append(fldSimple)

def add_title(text, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=12, bold=True):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_heading1(text, space_before=16, space_after=8):
    return add_title(text, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=space_before, space_after=space_after, bold=True)

def add_heading2(text, space_before=14, space_after=6):
    return add_title(text, size=14, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=space_before, space_after=space_after, bold=True)

def add_heading3(text, space_before=10, space_after=4):
    return add_title(text, size=13, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=space_before, space_after=space_after, bold=True)

def add_paragraph(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet(bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(bold_prefix)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.bold = True
    r1.font.color.rgb = RGBColor(15, 23, 42)

    r2 = p.add_run(text)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_code_block(code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    shading = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F8FAFC"/>')
    cell._tc.get_or_add_tcPr().append(shading)
    
    borders = parse_xml(r'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="6366F1"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(code_text)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(15, 23, 42)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)

def style_table(tbl, col_widths, headers, data):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    hdr_cells = tbl.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        hdr_cells[i].width = col_widths[i]
        shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="E2E8F0"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.space_before = Pt(3)
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            
    for row_idx, row_data in enumerate(data):
        row_cells = tbl.add_row().cells
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = str(text)
            row_cells[col_idx].width = col_widths[col_idx]
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx != 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(r'''
        <w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)

def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    return p