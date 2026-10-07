from report_helper import (
    doc, add_heading1, add_heading2,
    add_paragraph, add_bullet, add_caption
)
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def build_ch6_ch7_ch8():
    # ==================== CHAPTER 6: RESULTS AND SCREENSHOTS ====================
    add_heading1("CHAPTER 6\nRESULTS AND SCREENSHOTS", space_before=16, space_after=16)
    add_paragraph("The ExpenseIQ application was thoroughly tested and verified across all functional modules, covering CRUD operations, token-based authentication, database aggregation pipelines, and interactive Chart.js visualizations. The interface was evaluated under various data scenarios, including empty states, new registrations, category modifications, and multi-year financial trends.")
    add_paragraph("The following screenshots represent the live ExpenseIQ application interface, illustrating the system design and operational workflows.")

    if os.path.exists("report_assets/fig6_1_login.png"):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture("report_assets/fig6_1_login.png", width=Inches(5.5))
    add_caption("Figure 6.1: Authentication & Login Screen")

    if os.path.exists("report_assets/fig6_2_dashboard.png"):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture("report_assets/fig6_2_dashboard.png", width=Inches(5.5))
    add_caption("Figure 6.2: Financial Dashboard Screen (Summary Cards, Line & Doughnut Charts)")

    if os.path.exists("report_assets/fig6_3_transactions.png"):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture("report_assets/fig6_3_transactions.png", width=Inches(5.5))
    add_caption("Figure 6.3: Transactions Management & History Screen")

    doc.add_page_break()

    if os.path.exists("report_assets/fig6_4_categories.png"):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture("report_assets/fig6_4_categories.png", width=Inches(5.5))
    add_caption("Figure 6.4: Category Customization Screen (Icon & Color Picker)")

    if os.path.exists("report_assets/fig6_5_reports.png"):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture("report_assets/fig6_5_reports.png", width=Inches(5.5))
    add_caption("Figure 6.5: Monthly & Annual Financial Reports Screen")

    add_paragraph("All API endpoints, data filtering routines, and responsive layout workflows were validated as fully operational under rigorous testing conditions.")

    doc.add_page_break()

    # ==================== CHAPTER 7: CONCLUSION AND FUTURE ENHANCEMENT ====================
    add_heading1("CHAPTER 7\nCONCLUSION AND FUTURE ENHANCEMENT", space_before=16, space_after=16)
    add_paragraph("The ExpenseIQ: Personal Expense and Income Tracking System was successfully designed, implemented, and validated as a full-stack web application using the MERN stack. The system provides a centralized and interactive platform for managing incomes, expenses, categories, and monthly financial reports, eliminating the inefficiencies associated with manual bookkeeping and fragmented tracking methods.")
    add_paragraph("The project highlights the practical mastery of React 18 component design, Chart.js data visualization, Express.js REST API routing, JWT-based security middleware, Node.js asynchronous execution, and MongoDB aggregation pipelines. Controlled input validation and responsive dark glassmorphism styling ensure an engaging user experience.")
    add_paragraph("Through the development of ExpenseIQ, valuable hands-on competence was gained in full-stack architecture, RESTful API design, database schema modeling, client-side state management, and modern software engineering practices.")

    add_heading2("7.1 Future Enhancements")
    add_paragraph("The ExpenseIQ platform can be extended in future iterations with the following advanced capabilities:")
    add_bullet("Optical Character Recognition (OCR) Receipt Scanning: ", "Integrate machine learning or Tesseract.js OCR to automatically parse physical receipts and auto-populate expense fields.")
    add_bullet("Multi-Currency Support: ", "Integrate live foreign exchange rate APIs to support multi-currency logging and real-time conversion.")
    add_bullet("Budget Threshold Alerts: ", "Enable automated notifications and email alerts when monthly spending exceeds user-defined category limits.")
    add_bullet("Export to PDF & Excel: ", "Add one-click generation of formatted PDF statements and CSV/Excel expense spreadsheets.")
    add_bullet("Recurring Subscription Tracker: ", "Implement scheduled cron jobs to automatically track recurring bills, subscriptions, and salary credits.")

    doc.add_page_break()

    # ==================== CHAPTER 8: REFERENCES ====================
    add_heading1("CHAPTER 8\nREFERENCES", space_before=16, space_after=16)

    refs = [
        "1. React Documentation. \"Describing the UI & Managing State.\" Meta Platforms, Inc. Available: https://react.dev",
        "2. Express.js Guide. \"Routing, Middleware, and RESTful APIs.\" OpenJS Foundation. Available: https://expressjs.com",
        "3. MongoDB Documentation. \"The MongoDB Manual — Document Databases & Aggregation Pipelines.\" MongoDB, Inc. Available: https://www.mongodb.com/docs/",
        "4. Mongoose Documentation. \"Schemas, Models, Validation, and Population.\" Automattic. Available: https://mongoosejs.com/docs/",
        "5. Chart.js Documentation. \"Data Visualizations, Canvas Rendering, and Chart Types.\" Chart.js Open Source Project. Available: https://www.chartjs.org/docs/",
        "6. Fielding, Roy Thomas. \"Architectural Styles and the Design of Network-based Software Architectures.\" Ph.D. Dissertation, University of California, Irvine, 2000.",
        "7. Mozilla Developer Network (MDN). \"Fetch API, HTTP Status Codes, and Promises.\" Available: https://developer.mozilla.org",
        "8. Vite Documentation. \"Next Generation Frontend Tooling.\" Available: https://vitejs.dev",
        "9. Pressman, Roger S., and Bruce R. Maxim. \"Software Engineering: A Practitioner's Approach.\" 9th Edition, McGraw-Hill Education, 2020.",
        "10. W3C Recommendation. \"HTML5 & Cascading Style Sheets (CSS3) Specification.\" World Wide Web Consortium."
    ]

    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(ref)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(30, 41, 59)

    output_path = "c:/Users/SANJAY/.gemini/antigravity/playground/sidereal-halley/ExpenseIQ_Project_Report.docx"
    doc.save(output_path)
    print(f"SUCCESS: Report saved successfully to {output_path}")