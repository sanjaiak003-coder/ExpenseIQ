from report_helper import (
    doc, add_title, add_heading1, add_heading2, add_heading3,
    add_paragraph, add_bullet, style_table, add_caption
)
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

def build_preliminary_and_ch1_ch2():
    # ==================== PAGE 1: TITLE / COVER PAGE ====================
    add_title("JERUSALEM COLLEGE OF ENGINEERING", size=15, space_before=10, space_after=2)
    add_title("(An Autonomous Institution)", size=12, bold=False, space_before=0, space_after=2)
    add_title("(Approved by AICTE, Affiliated to Anna University, Chennai)", size=11, bold=False, space_before=0, space_after=2)
    add_title("ACCREDITED by NBA and NAAC with ‘A’ Grade", size=11, bold=True, space_before=0, space_after=2)
    add_title("Velachery Main Road, Narayanapuram, Pallikkaranai, Chennai – 600100", size=10.5, bold=False, space_before=0, space_after=14)

    if os.path.exists("report_assets/college_logo.png"):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(12)
        p_img.add_run().add_picture("report_assets/college_logo.png", width=Inches(1.2))

    add_title("MERN STACK DEVELOPMENT LABORATORY", size=14, space_before=6, space_after=2)
    add_title("JGE2542", size=13, space_before=0, space_after=2)
    add_title("[INDUSTRY SUPPORTED COURSE]", size=12, space_before=0, space_after=10)

    add_title("PROJECT REPORT", size=14, space_before=8, space_after=2)
    add_title("ACADEMIC YEAR 2026–27", size=11.5, bold=False, space_before=0, space_after=14)

    add_title("EXPENSEIQ: PERSONAL EXPENSE AND INCOME TRACKING SYSTEM", size=14, space_before=10, space_after=16)

    add_title("Submitted by", size=11, bold=False, space_before=8, space_after=2)
    add_title("JAYASOORIYA A", size=13, space_before=0, space_after=2)
    add_title("(2403310914821018)", size=12, space_before=0, space_after=8)

    add_title("III YEAR / V SEMESTER", size=12, space_before=4, space_after=2)
    add_title("REGULATION 2023", size=12, space_before=0, space_after=14)

    add_title("DEPARTMENT OF ARTIFICIAL INTELLIGENCE AND MACHINE LEARNING", size=13, space_before=10, space_after=2)
    add_title("B.E. CSE (AI & ML)", size=12, space_before=0, space_after=10)

    doc.add_page_break()

    # ==================== PAGE 2: BONAFIDE CERTIFICATE ====================
    add_heading1("BONAFIDE CERTIFICATE", space_before=20, space_after=24)
    add_paragraph("This is to certify that Mr. JAYASOORIYA A (2403310914821018), is a bonafide student of Jerusalem College of Engineering, Chennai. The student has successfully completed the project titled “EXPENSEIQ: PERSONAL EXPENSE AND INCOME TRACKING SYSTEM” as part of JGE2542 – MERN Stack Development Laboratory during the academic year 2026–2027. The project work has been completed successfully as per the requirements of the course.", space_after=40)

    tbl_sig = doc.add_table(rows=1, cols=2)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_s1, cell_s2 = tbl_sig.rows[0].cells
    cell_s1.width = Inches(3.0)
    cell_s2.width = Inches(3.0)

    p1 = cell_s1.paragraphs[0]
    p1.paragraph_format.line_spacing = 1.15
    r = p1.add_run("Ms. E. BRINDHA, M.E.\n")
    r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(12)
    r = p1.add_run("Faculty In-charge\nDept. of Artificial Intelligence and Machine Learning\nJerusalem College of Engineering\nPallikaranai, Chennai – 600100.")
    r.font.name = "Times New Roman"; r.font.size = Pt(11)

    p2 = cell_s2.paragraphs[0]
    p2.paragraph_format.line_spacing = 1.15
    r = p2.add_run("Dr. D. PARAMESWARI, M.Tech., Ph.D.\n")
    r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(12)
    r = p2.add_run("Professor and Head\nDept. of Artificial Intelligence and Machine Learning\nJerusalem College of Engineering\nPallikaranai, Chennai – 600100.")
    r.font.name = "Times New Roman"; r.font.size = Pt(11)

    p_viva = doc.add_paragraph()
    p_viva.paragraph_format.space_before = Pt(40)
    p_viva.paragraph_format.space_after = Pt(30)
    r = p_viva.add_run("Submitted for the Viva-Voce Examination held on ____________________")
    r.font.name = "Times New Roman"; r.font.size = Pt(12)

    tbl_ex = doc.add_table(rows=1, cols=2)
    tbl_ex.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_e1, cell_e2 = tbl_ex.rows[0].cells
    cell_e1.width = Inches(3.0)
    cell_e2.width = Inches(3.0)

    pe1 = cell_e1.paragraphs[0]
    pe1.paragraph_format.line_spacing = 1.15
    r = pe1.add_run("Internal Examiner\n")
    r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(12)
    r = pe1.add_run("Assistant Professor,\nJerusalem College of Engineering,\nChennai.")
    r.font.name = "Times New Roman"; r.font.size = Pt(11)

    pe2 = cell_e2.paragraphs[0]
    pe2.paragraph_format.line_spacing = 1.15
    r = pe2.add_run("External Examiner\n")
    r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(12)
    r = pe2.add_run("Co-Founder,\nHopeRaiser, Chennai.")
    r.font.name = "Times New Roman"; r.font.size = Pt(11)

    doc.add_page_break()

    # ==================== PAGE 3: ACKNOWLEDGEMENT ====================
    add_heading1("ACKNOWLEDGEMENT", space_before=20, space_after=20)
    add_paragraph("I express my sincere thanks to our honourable Chairperson Dr. M. MALA, M.A., M.Phil., and our Principal, Dr. S. SATHIYAMURTHY, M.E., Ph.D., for providing me all the facilities and resources required to complete my project.")
    add_paragraph("I am deeply grateful to Dr. D. PARAMESWARI, Professor and Head, Department of Artificial Intelligence and Machine Learning, for being a constant source of inspiration throughout the course of my project.")
    add_paragraph("I also express my sincere thanks to the faculties of the Department of R&D, Jerusalem College of Engineering, for conducting the Industry Supported Course on MERN Stack Development for the students in association with HopeRaiser.")
    add_paragraph("I express my sincere gratitude to Mr. B. Timo Jacob and Ms. R. Keerthana, Co-Founders of HopeRaiser, Chennai, who taught this Industry Supported Course to the students, and for serving as the External Industry Experts for my project viva-voce. Their valuable inputs, technical insights, and constructive evaluation greatly helped me improve my project outcomes.")
    add_paragraph("I express my deep sense of gratitude to all faculties of our department for their guidance and useful suggestions, which helped me in completing this project.")
    add_paragraph("Finally, I thank the almighty, my parents and friends for their constant encouragement, without which this project would not have been possible.", space_after=40)

    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_sign.add_run("JAYASOORIYA A")
    r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(13)

    doc.add_page_break()

    # ==================== PAGE 4: ABSTRACT ====================
    add_heading1("ABSTRACT", space_before=20, space_after=20)
    add_paragraph("Managing personal finances, budgeting daily expenditures, and analyzing recurring cash flows represent essential life skills in contemporary society. However, most individuals continue to record transactions through fragmented paper diaries, ad-hoc spreadsheets, or disconnected notes. This manual paradigm leads to calculation discrepancies, lack of visual insight, poor budgeting discipline, and absence of historical financial analytics.")
    add_paragraph("ExpenseIQ is a modern, responsive, and full-featured Personal Expense and Income Tracking System engineered using the MERN stack (MongoDB, Express.js, React.js, and Node.js). The application manages three core collegiate and personal finance domains: User Authentication & Profile Management, Category Customization, and Transaction Logging & Monthly Analytics. It provides a robust RESTful API featuring structured JSON endpoints, standard HTTP status handling, and token-based JSON Web Token (JWT) security.")
    add_paragraph("The client-side architecture leverages React 18, React Router DOM, and Chart.js alongside a custom dark glassmorphism design system to deliver an intuitive single-page application (SPA). Real-time visual summaries, interactive monthly trend lines, category distribution doughnut charts, and comprehensive 12-month report tables enable users to make data-driven budgeting decisions. The project successfully demonstrates the implementation of full-stack JavaScript technologies, unidirectional state management, and NoSQL aggregation pipelines.")

    doc.add_page_break()

    # ==================== PAGE 5: TABLE OF CONTENTS ====================
    add_heading1("TABLE OF CONTENTS", space_before=16, space_after=18)
    toc_headers = ["S.No", "Chapter / Content", "Page No"]
    toc_col_widths = [Inches(0.8), Inches(4.5), Inches(1.0)]
    toc_data = [
        ["1", "Introduction", "1"],
        ["2", "  1.1 Project Overview", "1"],
        ["3", "Requirements Analysis", "2"],
        ["4", "  2.1 Hardware Requirements", "2"],
        ["5", "  2.2 Software Requirements", "2"],
        ["6", "  2.3 Functional Requirements", "3"],
        ["7", "  2.4 Non-Functional Requirements", "3"],
        ["8", "System Design", "4"],
        ["9", "  3.1 System Architecture", "4"],
        ["10", "  3.2 Database Design", "6"],
        ["11", "Technology Stack", "8"],
        ["12", "  4.1 MongoDB", "8"],
        ["13", "  4.2 Express.js", "8"],
        ["14", "  4.3 React.js", "8"],
        ["15", "  4.4 Node.js", "8"],
        ["16", "  4.5 Other Tools / Libraries", "9"],
        ["17", "System Implementation", "10"],
        ["18", "  5.1 Module 1 – Authentication & User Management", "10"],
        ["19", "  5.2 Module 2 – Category & Transaction Management", "10"],
        ["20", "  5.3 API / Backend Implementation - code", "11"],
        ["21", "  5.4 Frontend Implementation - code", "13"],
        ["22", "Results and Screenshots", "14"],
        ["23", "Conclusion and Future Enhancement", "17"],
        ["24", "References", "18"]
    ]
    tbl_toc = doc.add_table(rows=1, cols=3)
    style_table(tbl_toc, toc_col_widths, toc_headers, toc_data)

    doc.add_page_break()

    # ==================== CHAPTER 1: INTRODUCTION ====================
    add_heading1("CHAPTER 1\nINTRODUCTION", space_before=16, space_after=16)
    add_heading2("1.1 Project Overview")
    add_paragraph("ExpenseIQ: Personal Expense and Income Tracking System is a full-stack web application designed to simplify, automate, and centralize personal budgeting, income monitoring, and expenditure tracking. In traditional setups, individuals record financial transactions through paper notebooks, manual spreadsheets, or isolated calculator apps. Managing finances across disparate tools frequently results in arithmetic errors, forgotten expenses, zero category-level insights, and difficulty maintaining long-term financial health. ExpenseIQ provides a unified, web-based platform to solve these challenges.")
    add_paragraph("The application provides comprehensive modules for recording incomes and expenses, organizing custom color-coded categories, generating real-time balance calculations, and visualizing spending patterns through interactive charts. Users can filter transaction histories by type, month, or year, manage custom categories with emoji icons, and inspect annual net savings rates with instant visual feedback.")
    add_paragraph("ExpenseIQ is developed using the MERN stack, comprising MongoDB, Express.js, React.js, and Node.js. React.js with Vite is utilized to build the single-page application frontend, styled with a modern glassmorphism dark theme. Node.js and Express.js power the backend REST API, handling secure authentication and database operations. MongoDB stores structured records for users, categories, and transactions, while Mongoose provides strict schema validation and efficient aggregation pipelines.")
    add_paragraph("The system follows a three-tier client-server architecture, where the React client communicates securely with the Express backend using JSON Web Tokens (JWT) and Axios HTTP requests. The backend executes full CRUD (Create, Read, Update, Delete) operations, ensuring real-time data synchronization between the UI and the database.")
    add_paragraph("To deliver a superior user experience, the system incorporates responsive dashboard analytics, modal workflows for seamless additions/edits, automated default category seeding, and graceful offline fallback handling. Overall, ExpenseIQ demonstrates modern full-stack development best practices including modular component design, RESTful service architecture, NoSQL aggregation, and stateful React interfaces.")

    doc.add_page_break()

    # ==================== CHAPTER 2: REQUIREMENTS ANALYSIS ====================
    add_heading1("CHAPTER 2\nREQUIREMENTS ANALYSIS", space_before=16, space_after=16)
    add_heading2("2.1 Hardware Requirements")
    add_paragraph("The hardware configuration required to develop and execute the ExpenseIQ application is listed in Table 2.1.")
    tbl_hw = doc.add_table(rows=1, cols=3)
    hw_headers = ["Component", "Minimum", "Recommended"]
    hw_widths = [Inches(1.8), Inches(2.2), Inches(2.3)]
    hw_data = [
        ["Processor", "Intel Core i3 or equivalent", "Intel Core i5 / AMD Ryzen 5 or higher"],
        ["RAM", "4 GB", "8 GB or more"],
        ["Storage", "10 GB free disk space", "20 GB or more (SSD)"],
        ["Display", "1366 × 768 resolution", "1920 × 1080 resolution"],
        ["Keyboard", "Standard keyboard", "Standard keyboard"],
        ["Mouse", "Standard mouse", "Optical mouse"],
        ["Internet connection", "Required for package installation", "Broadband connection"]
    ]
    style_table(tbl_hw, hw_widths, hw_headers, hw_data)
    add_caption("Table 2.1: Hardware Requirements")

    add_heading2("2.2 Software Requirements")
    add_paragraph("The software environment and development tooling used for building and hosting the application are detailed in Table 2.2.")
    tbl_sw = doc.add_table(rows=1, cols=2)
    sw_headers = ["Software", "Details"]
    sw_widths = [Inches(2.3), Inches(4.0)]
    sw_data = [
        ["Operating System", "Microsoft Windows 10/11, macOS, or Linux"],
        ["Visual Studio Code", "Source code editor used for frontend and backend development"],
        ["Node.js (v16+)", "JavaScript runtime used for backend server execution and npm tools"],
        ["MongoDB / Atlas", "NoSQL document database used to store users, categories, and transactions"],
        ["MongoDB Compass", "Graphical desktop interface for inspecting and managing MongoDB collections"],
        ["Web Browser", "Google Chrome, Microsoft Edge, or Mozilla Firefox"],
        ["Git / GitHub", "Distributed version control system and repository hosting"]
    ]
    style_table(tbl_sw, sw_widths, sw_headers, sw_data)
    add_caption("Table 2.2: Software Requirements")

    add_heading2("2.3 Functional Requirements")
    add_paragraph("The core functional capabilities of the ExpenseIQ system are specified in Table 2.3.")
    tbl_fr = doc.add_table(rows=1, cols=3)
    fr_headers = ["ID", "Requirement", "Description"]
    fr_widths = [Inches(0.8), Inches(1.8), Inches(3.7)]
    fr_data = [
        ["FR1", "User Authentication", "Register, login, and secure user sessions using JWT and bcryptjs hashing."],
        ["FR2", "Category Management", "Create, customize (color, icon), view, update, and delete income/expense categories."],
        ["FR3", "Transaction Tracking", "Log income and expense transactions with amount, category, date, and description."],
        ["FR4", "Transaction Filtering", "Filter transaction history by type, category, month, and year with pagination."],
        ["FR5", "Interactive Dashboard", "Display real-time summary cards, monthly trend lines, and category doughnut charts."],
        ["FR6", "Monthly & Annual Reports", "Generate aggregated 12-month reports, savings rates, and category progress bars."],
        ["FR7", "RESTful API Endpoints", "Provide structured endpoints using GET, POST, PUT, and DELETE HTTP verbs."]
    ]
    style_table(tbl_fr, fr_widths, fr_headers, fr_data)
    add_caption("Table 2.3: Functional Requirements")

    add_heading2("2.4 Non-Functional Requirements")
    add_paragraph("The operational quality attributes and performance parameters of the system are outlined in Table 2.4.")
    tbl_nfr = doc.add_table(rows=1, cols=2)
    nfr_headers = ["Attribute", "Requirement"]
    nfr_widths = [Inches(2.0), Inches(4.3)]
    nfr_data = [
        ["Usability", "Modern dark glassmorphism user interface that is intuitive and easy to navigate."],
        ["Performance", "Fast REST API response times (<100ms locally) and lightweight bundle sizes."],
        ["Security", "Protected routes, password encryption with bcrypt, and token expiration handling."],
        ["Robustness", "Graceful error handling, input validation, and database connection fallback modes."],
        ["Maintainability", "Clean separation of routes, models, middleware, and reusable React components."],
        ["Compatibility", "Cross-browser support for modern web browsers and responsive mobile screen layouts."]
    ]
    style_table(tbl_nfr, nfr_widths, nfr_headers, nfr_data)
    add_caption("Table 2.4: Non-Functional Requirements")

    doc.add_page_break()